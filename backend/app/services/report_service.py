import io
import json
from datetime import datetime, timezone
from typing import Optional, List
from sqlalchemy.orm import Session

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable

from app.models import Decision, Approval, ApprovalHistory, AuditLog, User, Alternative


# --- Colors & Styles ---
NAVY_HEADER = colors.HexColor("#1e293b")
ACCENT_BLUE = colors.HexColor("#2563eb")
LIGHT_BG = colors.HexColor("#f8fafc")
BORDER_COLOR = colors.HexColor("#cbd5e1")


def _get_pdf_styles():
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=18,
        leading=22,
        textColor=NAVY_HEADER,
        spaceAfter=10
    )
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#64748b"),
        spaceAfter=15
    )
    heading2_style = ParagraphStyle(
        'SectionHeading',
        parent=styles['Heading2'],
        fontSize=13,
        leading=16,
        textColor=ACCENT_BLUE,
        spaceBefore=12,
        spaceAfter=6
    )
    body_style = ParagraphStyle(
        'BodyTextCustom',
        parent=styles['Normal'],
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#334155")
    )
    table_cell_style = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontSize=8,
        leading=11,
        textColor=colors.HexColor("#1e293b")
    )
    table_header_style = ParagraphStyle(
        'TableHeader',
        parent=styles['Normal'],
        fontSize=8,
        leading=11,
        textColor=colors.white,
        fontName='Helvetica-Bold'
    )
    return {
        'title': title_style,
        'subtitle': subtitle_style,
        'heading2': heading2_style,
        'body': body_style,
        'cell': table_cell_style,
        'header': table_header_style
    }


# =========================================================================
# 1. DECISION SUMMARY REPORT
# =========================================================================

def generate_decisions_pdf(db: Session, decision_id: Optional[int] = None) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
    styles = _get_pdf_styles()
    story = []

    query = db.query(Decision)
    if decision_id:
        query = query.filter(Decision.id == decision_id)
    decisions = query.order_by(Decision.id.desc()).all()

    # Title Banner
    title_text = f"Decision Replay Platform - {'Decision Dossier' if decision_id else 'Organizational Decisions Summary'}"
    story.append(Paragraph(title_text, styles['title']))
    story.append(Paragraph(f"Generated on {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')} • Total Decisions: {len(decisions)}", styles['subtitle']))
    story.append(HRFlowable(width="100%", thickness=1, color=BORDER_COLOR, spaceAfter=15))

    for dec in decisions:
        story.append(Paragraph(f"#{dec.id} - {dec.title}", styles['heading2']))
        
        meta_data = [
            [
                Paragraph("<b>Category:</b> " + (dec.category or "N/A"), styles['cell']),
                Paragraph("<b>Status:</b> " + dec.status.value, styles['cell']),
                Paragraph("<b>Author:</b> " + (dec.creator.full_name if dec.creator else "Unknown"), styles['cell']),
                Paragraph("<b>Date:</b> " + dec.created_at.strftime('%Y-%m-%d'), styles['cell'])
            ]
        ]
        meta_table = Table(meta_data, colWidths=[130, 130, 140, 140])
        meta_table.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,-1), LIGHT_BG),
            ('BOX', (0,0), (-1,-1), 0.5, BORDER_COLOR),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('TOPPADDING', (0,0), (-1,-1), 4),
            ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ]))
        story.append(meta_table)
        story.append(Spacer(1, 6))

        # Problem Statement
        story.append(Paragraph("<b>Problem Statement:</b>", styles['body']))
        story.append(Paragraph(dec.problem_statement or "No problem statement recorded.", styles['body']))
        story.append(Spacer(1, 8))

        # Alternatives Table
        if dec.alternatives:
            story.append(Paragraph("<b>Evaluated Alternatives:</b>", styles['body']))
            alt_rows = [[
                Paragraph("Title", styles['header']),
                Paragraph("Cost", styles['header']),
                Paragraph("Feasibility", styles['header']),
                Paragraph("Pros / Trade-offs", styles['header']),
                Paragraph("Risk Assessment", styles['header'])
            ]]
            for alt in dec.alternatives:
                pros_str = ", ".join(alt.parsed_pros[:3]) if alt.parsed_pros else "N/A"
                alt_rows.append([
                    Paragraph(f"<b>{alt.title}</b>", styles['cell']),
                    Paragraph(f"${alt.estimated_cost:,.2f}" if alt.estimated_cost else "$0.00", styles['cell']),
                    Paragraph(f"{alt.feasibility_score}/10", styles['cell']),
                    Paragraph(pros_str, styles['cell']),
                    Paragraph(alt.risk_assessment or "N/A", styles['cell'])
                ])
            alt_table = Table(alt_rows, colWidths=[120, 60, 60, 150, 150])
            alt_table.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), ACCENT_BLUE),
                ('BOX', (0,0), (-1,-1), 0.5, BORDER_COLOR),
                ('GRID', (0,0), (-1,-1), 0.5, BORDER_COLOR),
                ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, LIGHT_BG]),
                ('VALIGN', (0,0), (-1,-1), 'TOP'),
                ('TOPPADDING', (0,0), (-1,-1), 3),
                ('BOTTOMPADDING', (0,0), (-1,-1), 3),
            ]))
            story.append(alt_table)
            story.append(Spacer(1, 8))

        # Approval Trail
        if dec.approvals:
            story.append(Paragraph("<b>Approval Lifecycle Chain:</b>", styles['body']))
            app_rows = [[
                Paragraph("Level", styles['header']),
                Paragraph("Approver", styles['header']),
                Paragraph("Status", styles['header']),
                Paragraph("Comments / Feedback", styles['header']),
                Paragraph("Updated At", styles['header'])
            ]]
            for app in dec.approvals:
                app_rows.append([
                    Paragraph(f"Level {app.level}", styles['cell']),
                    Paragraph(app.approver.full_name if app.approver else "Unassigned", styles['cell']),
                    Paragraph(app.status.value, styles['cell']),
                    Paragraph(app.comments or "No comments", styles['cell']),
                    Paragraph(app.updated_at.strftime('%Y-%m-%d %H:%M'), styles['cell'])
                ])
            app_table = Table(app_rows, colWidths=[50, 130, 80, 180, 100])
            app_table.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), NAVY_HEADER),
                ('BOX', (0,0), (-1,-1), 0.5, BORDER_COLOR),
                ('GRID', (0,0), (-1,-1), 0.5, BORDER_COLOR),
                ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, LIGHT_BG]),
                ('VALIGN', (0,0), (-1,-1), 'TOP'),
                ('TOPPADDING', (0,0), (-1,-1), 3),
                ('BOTTOMPADDING', (0,0), (-1,-1), 3),
            ]))
            story.append(app_table)
            story.append(Spacer(1, 8))

        story.append(HRFlowable(width="100%", thickness=0.5, color=BORDER_COLOR, spaceAfter=14))

    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()


def generate_decisions_excel(db: Session, decision_id: Optional[int] = None) -> bytes:
    wb = Workbook()
    
    # Sheet 1: Decisions
    ws_dec = wb.active
    ws_dec.title = "Decisions Summary"
    
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=10)
    center_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    headers = ["ID", "Title", "Category", "Status", "Creator", "Created At", "Problem Statement", "Alternatives Count"]
    ws_dec.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws_dec.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    query = db.query(Decision)
    if decision_id:
        query = query.filter(Decision.id == decision_id)
    decisions = query.order_by(Decision.id.desc()).all()

    for row_idx, dec in enumerate(decisions, start=2):
        ws_dec.append([
            dec.id,
            dec.title,
            dec.category,
            dec.status.value,
            dec.creator.full_name if dec.creator else "Unknown",
            dec.created_at.strftime('%Y-%m-%d %H:%M'),
            dec.problem_statement,
            len(dec.alternatives)
        ])
        for col_idx in range(1, len(headers) + 1):
            ws_dec.cell(row=row_idx, column=col_idx).border = thin_border

    # Sheet 2: Alternatives
    ws_alt = wb.create_sheet(title="Evaluated Alternatives")
    alt_headers = ["Decision ID", "Decision Title", "Alternative Title", "Estimated Cost ($)", "Feasibility (1-10)", "Pros", "Cons", "Risk Assessment"]
    ws_alt.append(alt_headers)
    for col_num in range(1, len(alt_headers) + 1):
        cell = ws_alt.cell(row=1, column=col_num)
        cell.fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
        cell.font = header_font
        cell.alignment = center_align

    alt_row_idx = 2
    for dec in decisions:
        for alt in dec.alternatives:
            ws_alt.append([
                dec.id,
                dec.title,
                alt.title,
                alt.estimated_cost or 0.0,
                alt.feasibility_score or 5,
                ", ".join(alt.parsed_pros),
                ", ".join(alt.parsed_cons),
                alt.risk_assessment or ""
            ])
            for col_idx in range(1, len(alt_headers) + 1):
                ws_alt.cell(row=alt_row_idx, column=col_idx).border = thin_border
            alt_row_idx += 1

    # Auto-adjust column widths
    for sheet in [ws_dec, ws_alt]:
        for col in sheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

    out_buffer = io.BytesIO()
    wb.save(out_buffer)
    out_buffer.seek(0)
    return out_buffer.getvalue()


# =========================================================================
# 2. APPROVAL TURNAROUND & TEAM PRODUCTIVITY REPORT
# =========================================================================

def generate_approvals_pdf(db: Session) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
    styles = _get_pdf_styles()
    story = []

    story.append(Paragraph("Approval Turnaround & Team Productivity Report", styles['title']))
    story.append(Paragraph(f"Generated on {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')} • Enterprise SLA & Governance", styles['subtitle']))
    story.append(HRFlowable(width="100%", thickness=1, color=BORDER_COLOR, spaceAfter=15))

    approvals = db.query(Approval).order_by(Approval.created_at.desc()).all()

    # Summary Metrics
    total_count = len(approvals)
    approved_count = sum(1 for a in approvals if a.status.value == "APPROVED")
    pending_count = sum(1 for a in approvals if a.status.value == "PENDING")
    rejected_count = sum(1 for a in approvals if a.status.value == "REJECTED")

    summary_data = [
        [
            Paragraph(f"<b>Total Approvals:</b> {total_count}", styles['cell']),
            Paragraph(f"<b>Approved:</b> {approved_count}", styles['cell']),
            Paragraph(f"<b>Pending Queue:</b> {pending_count}", styles['cell']),
            Paragraph(f"<b>Rejected:</b> {rejected_count}", styles['cell'])
        ]
    ]
    sum_table = Table(summary_data, colWidths=[135, 135, 135, 135])
    sum_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), LIGHT_BG),
        ('BOX', (0,0), (-1,-1), 1, BORDER_COLOR),
        ('TOPPADDING', (0,0), (-1,-1), 6),
        ('BOTTOMPADDING', (0,0), (-1,-1), 6),
    ]))
    story.append(sum_table)
    story.append(Spacer(1, 14))

    # Approvals Table
    app_rows = [[
        Paragraph("ID", styles['header']),
        Paragraph("Decision", styles['header']),
        Paragraph("Level", styles['header']),
        Paragraph("Approver", styles['header']),
        Paragraph("Status", styles['header']),
        Paragraph("Comments / Feedback", styles['header']),
        Paragraph("Created At", styles['header'])
    ]]
    for app in approvals:
        app_rows.append([
            Paragraph(str(app.id), styles['cell']),
            Paragraph(app.decision.title if app.decision else "N/A", styles['cell']),
            Paragraph(f"L{app.level}", styles['cell']),
            Paragraph(app.approver.full_name if app.approver else "Unassigned", styles['cell']),
            Paragraph(app.status.value, styles['cell']),
            Paragraph(app.comments or "-", styles['cell']),
            Paragraph(app.created_at.strftime('%Y-%m-%d %H:%M'), styles['cell'])
        ])

    table = Table(app_rows, colWidths=[30, 150, 40, 100, 70, 90, 60])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), NAVY_HEADER),
        ('BOX', (0,0), (-1,-1), 0.5, BORDER_COLOR),
        ('GRID', (0,0), (-1,-1), 0.5, BORDER_COLOR),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, LIGHT_BG]),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('TOPPADDING', (0,0), (-1,-1), 3),
        ('BOTTOMPADDING', (0,0), (-1,-1), 3),
    ]))
    story.append(table)

    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()


def generate_approvals_excel(db: Session) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Approval Turnaround"

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=10)
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    headers = ["Approval ID", "Decision ID", "Decision Title", "Level", "Approver Name", "Status", "Comments", "Created At", "Updated At"]
    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font

    approvals = db.query(Approval).order_by(Approval.created_at.desc()).all()
    for row_idx, app in enumerate(approvals, start=2):
        ws.append([
            app.id,
            app.decision_id,
            app.decision.title if app.decision else "",
            f"Level {app.level}",
            app.approver.full_name if app.approver else "",
            app.status.value,
            app.comments or "",
            app.created_at.strftime('%Y-%m-%d %H:%M:%S'),
            app.updated_at.strftime('%Y-%m-%d %H:%M:%S')
        ])
        for col_idx in range(1, len(headers) + 1):
            ws.cell(row=row_idx, column=col_idx).border = thin_border

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    out_buffer = io.BytesIO()
    wb.save(out_buffer)
    out_buffer.seek(0)
    return out_buffer.getvalue()


# =========================================================================
# 3. AUDIT TRAIL & COMPLIANCE REPORT
# =========================================================================

def generate_audit_pdf(db: Session) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
    styles = _get_pdf_styles()
    story = []

    story.append(Paragraph("Audit Trail & Compliance Report", styles['title']))
    story.append(Paragraph(f"Generated on {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')} • Enterprise Security & Access Tracking", styles['subtitle']))
    story.append(HRFlowable(width="100%", thickness=1, color=BORDER_COLOR, spaceAfter=15))

    logs = db.query(AuditLog).order_by(AuditLog.timestamp.desc()).limit(200).all()

    rows = [[
        Paragraph("Time (UTC)", styles['header']),
        Paragraph("Action", styles['header']),
        Paragraph("Resource", styles['header']),
        Paragraph("Target ID", styles['header']),
        Paragraph("User", styles['header']),
        Paragraph("IP Address", styles['header']),
        Paragraph("Audit Event Details", styles['header'])
    ]]
    for l in logs:
        details_str = json.dumps(l.parsed_details) if l.parsed_details else "-"
        rows.append([
            Paragraph(l.timestamp.strftime('%Y-%m-%d %H:%M'), styles['cell']),
            Paragraph(l.action, styles['cell']),
            Paragraph(l.resource_type, styles['cell']),
            Paragraph(l.resource_id or "-", styles['cell']),
            Paragraph(l.user.full_name if l.user else "System", styles['cell']),
            Paragraph(l.ip_address or "Local", styles['cell']),
            Paragraph(details_str[:80] + ("..." if len(details_str) > 80 else ""), styles['cell'])
        ])

    table = Table(rows, colWidths=[70, 55, 60, 45, 90, 60, 160])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), NAVY_HEADER),
        ('BOX', (0,0), (-1,-1), 0.5, BORDER_COLOR),
        ('GRID', (0,0), (-1,-1), 0.5, BORDER_COLOR),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, LIGHT_BG]),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('TOPPADDING', (0,0), (-1,-1), 3),
        ('BOTTOMPADDING', (0,0), (-1,-1), 3),
    ]))
    story.append(table)

    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()


def generate_audit_excel(db: Session) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Audit Trail"

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=10)
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    headers = ["Audit ID", "Timestamp (UTC)", "Action", "Resource Type", "Resource ID", "User ID", "User Name", "IP Address", "Details"]
    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font

    logs = db.query(AuditLog).order_by(AuditLog.timestamp.desc()).all()
    for row_idx, l in enumerate(logs, start=2):
        ws.append([
            l.id,
            l.timestamp.strftime('%Y-%m-%d %H:%M:%S'),
            l.action,
            l.resource_type,
            l.resource_id or "",
            l.user_id or "",
            l.user.full_name if l.user else "System",
            l.ip_address or "",
            json.dumps(l.parsed_details) if l.parsed_details else ""
        ])
        for col_idx in range(1, len(headers) + 1):
            ws.cell(row=row_idx, column=col_idx).border = thin_border

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    out_buffer = io.BytesIO()
    wb.save(out_buffer)
    out_buffer.seek(0)
    return out_buffer.getvalue()
